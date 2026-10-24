from pathlib import Path
import re,json,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parents[1]
names=['Final scientific report','Scientific manuscript','Cox proportional hazards analysis','README']
breaks={
'Final scientific report':{'Research questions and design','How the analysis data were prepared','Final descriptive findings','Timing and clinical interpretation','Statistical methodology and exploratory associations','Observation sensitivity and the bleeding component','Limitations and final conclusions','References'},
'Scientific manuscript':{'Introduction','Statistical analysis','Results continued','Discussion','Tables and figure legends','References'},
'Cox proportional hazards analysis':{'Colonoscopy hazard ratios','EGD hazard ratios','Proportional hazards diagnostics and model limits','Early and late EGD associations'},
'README':set()}
widths={
'Final scientific report':[[1.88,5.37],[1.74,5.51],[1.74,5.51],[.90,.90,1.80,1.80,1.85],[1.11,1.6,2.22,2.32],[3.75,1.04,.90,1.56],[4.3,.90,.90,1.15]],
}
def inline(p,s):
 for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)',s):
  if token.startswith('**') and token.endswith('**'):p.add_run(token[2:-2]).bold=True
  elif token.startswith('`') and token.endswith('`'):
   r=p.add_run(token[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
  else:p.add_run(token)
def plain(s):return re.sub(r'\*\*|`','',s)
audit=[]
for name in names:
 text=(R/(name+'.md')).read_text(encoding='utf-8')
 if name=='README':
  text=text.replace('**Final scientific report.pdf**','**Final scientific report.docx**').replace('**Scientific manuscript.pdf** and its matching Markdown file','**Scientific manuscript.docx**').replace('**Cox proportional hazards analysis.pdf**','**Cox proportional hazards analysis.docx**')
  text=text.replace('## Files','The four Word files contain editable text and tables. The scientific report includes figures as images; original SVG and PNG figure files remain in the figures folder. PDFs are retained as fixed-layout reference copies.\n\n## Files')
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=Inches(.6);sec.bottom_margin=Inches(.6);sec.left_margin=Inches(.625);sec.right_margin=Inches(.625)
 for st in ['Normal','Title','Heading 1','Heading 2']:
  style=d.styles[st];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.10;d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.styles['Title'].font.size=Pt(21);d.styles['Title'].font.bold=True
 d.styles['Heading 1'].font.size=Pt(14);d.styles['Heading 1'].paragraph_format.space_before=Pt(8);d.styles['Heading 1'].paragraph_format.space_after=Pt(8)
 d.styles['Heading 1'].paragraph_format.keep_with_next=True
 footer=sec.footer.paragraphs[0];footer.alignment=2;r=footer.add_run();r.font.size=Pt(8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
 lines=text.splitlines();i=0;tablecount=0;expected=[]
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('```'):
   i+=1
   while i<len(lines) and not lines[i].startswith('```'):
    p=d.add_paragraph();r=p.add_run(lines[i]);r.font.name='Consolas';r.font.size=Pt(8);expected.append(lines[i]);i+=1
   i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    vals=[s.strip() for s in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',s) for s in vals):rows.append(vals)
    i+=1
   t=d.add_table(rows=1,cols=len(rows[0]));t.autofit=False
   # Widths depend on visible header semantics, not table ordinal.
   n=len(rows[0]); ws=[7.25/n]*n
   if n==2:ws=[2.0,5.25]
   if n==3:ws=[1.5,2.875,2.875]
   if n==4:ws=[3.75,1.05,.90,1.55] if 'Cohort and' in rows[0][0] else ([4.25,1.,1.,1.] if 'Follow-up' in rows[0][0] else [1.8,1.82,1.82,1.81])
   if n==5:ws=[1.1,1.54,1.54,1.54,1.53] if rows[0][0]=='Predictor' else [.9,.9,1.8,1.8,1.85]
   for col,w in zip(t.columns,ws):col.width=Inches(w)
   for ri,row in enumerate(rows):
    cells=t.rows[0].cells if ri==0 else t.add_row().cells
    for cell,w,value in zip(cells,ws,row):
     cell.width=Inches(w);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
     inline(p,value);expected.append(plain(value))
     for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
     cell.vertical_alignment=0
     if ri==0:
      sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF1');cell._tc.get_or_add_tcPr().append(sh)
    trpr=t.rows[ri]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    if ri==0:trpr.append(OxmlElement('w:tblHeader'))
   d.add_paragraph().paragraph_format.space_after=Pt(2);tablecount+=1;continue
  if line.startswith('!['):
   m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
   if not m:raise ValueError(line)
   d.add_picture(str(R/m.group(2)),width=Inches(7.0));p=d.add_paragraph(m.group(1));p.style='Caption';p.paragraph_format.keep_with_next=False
   expected.append(m.group(1));i+=1;continue
  if line.startswith('# '):
   s=line[2:];d.add_paragraph(s,'Title');expected.append(s)
  elif line.startswith('## '):
   s=line[3:];p=d.add_paragraph(s,'Heading 1');p.paragraph_format.page_break_before=s in breaks[name];expected.append(s)
  else:
   bullet=line.startswith('- ');p=d.add_paragraph(style='List Bullet' if bullet else 'Normal');s=line[2:] if bullet else line;inline(p,s);expected.append(plain(s))
  i+=1
 d.core_properties.title=d.paragraphs[0].text;d.core_properties.author='Research project team'
 path=R/(name+'.docx');d.save(path)
 # Validate paragraph/table content in document order and confirm no editing protection.
 from docx.text.paragraph import Paragraph
 from docx.table import Table
 actual=[]
 for node in d.element.body:
  if node.tag==qn('w:p'):
   p=Paragraph(node,d)
   if p.text:actual.append(p.text)
  elif node.tag==qn('w:tbl'):
   for row in Table(node,d).rows:
    for cell in row.cells:actual.append(cell.text)
 assert actual==expected,(name,'Content mismatch')
 with zipfile.ZipFile(path) as z:
  assert b'documentProtection' not in z.read('word/settings.xml')
 audit.append({'file':path.name,'editable_tables':tablecount,'figures':len(d.inline_shapes),'text_blocks_verified':len(expected),'all_content_matched':True,'editing_protected':False})
(R/'documentation/word_conversion_validation.json').write_text(json.dumps(audit,indent=2))
print(json.dumps(audit,indent=2))
